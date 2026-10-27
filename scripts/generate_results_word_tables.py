from __future__ import annotations

import csv
import math
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Callable, Iterable, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
OUTPUT = RESULTS / "核心实验结果数据表.docx"


COLUMN_LABELS = {
    "method": "方法",
    "variant": "实验方案",
    "seed": "随机种子",
    "seeds": "种子数",
    "completed_seeds": "完成种子数",
    "carbon_price": "碳价",
    "preference": "偏好",
    "season": "季节",
    "metric": "指标",
    "comparison": "对比项",
    "comparator": "对比方法",
    "ablation": "消融编号",
    "ablation_method": "消融方法",
    "market": "交易机制",
    "algorithm": "优化算法",
    "effect": "作用项",
    "hv_mean": "HV均值",
    "hv_std": "HV标准差",
    "igd_mean": "IGD均值",
    "igd_std": "IGD标准差",
    "sparsity_mean": "稀疏度均值",
    "sparsity_std": "稀疏度标准差",
    "front_points_mean": "帕累托点数均值",
    "front_points_std": "帕累托点数标准差",
    "front_points": "帕累托点数",
    "raw_hypervolume": "原始超体积",
    "normalized_hv": "归一化超体积",
    "normalized_igd": "归一化IGD",
    "normalized_sparsity": "归一化稀疏度",
    "min_cost_mean": "最低成本均值",
    "min_carbon_mean": "最低碳排放均值",
    "min_annual_cost": "最低全年成本/CNY",
    "max_annual_cost": "最高全年成本/CNY",
    "mean_annual_cost": "平均全年成本/CNY",
    "min_annual_carbon": "最低全年碳排放/kgCO2",
    "max_annual_carbon": "最高全年碳排放/kgCO2",
    "mean_annual_carbon": "平均全年碳排放/kgCO2",
    "hv_reference_cost": "HV成本参考点/CNY",
    "hv_reference_carbon": "HV碳排放参考点/kgCO2",
    "wins": "胜出数",
    "ties": "持平数",
    "losses": "落后数",
    "main_wins": "主方法胜出数",
    "main_losses": "主方法落后数",
    "mean_improvement": "平均改善量",
    "one_sided_p": "单侧检验p值",
    "one_sided_wilcoxon_p": "单侧Wilcoxon p值",
    "one_sided_wilcoxon_pvalue": "单侧Wilcoxon p值",
    "cohen_dz": "Cohen dz",
    "bootstrap_ci_low": "Bootstrap区间下限",
    "bootstrap_ci_high": "Bootstrap区间上限",
    "evaluations": "评估次数",
    "hours_mean": "耗时均值/h",
    "points_mean": "帕累托点数均值",
    "points_std": "帕累托点数标准差",
    "hv_auc_mean": "HV-AUC均值",
    "hv_auc_std": "HV-AUC标准差",
    "igd_auc_mean": "IGD-AUC均值",
    "igd_auc_std": "IGD-AUC标准差",
    "hv_change_mean_percent": "HV平均变化/%",
    "hv_change_std_percent": "HV变化标准差/%",
    "hv_improved_seeds": "HV改善种子数",
    "igd_reduction_mean_percent": "IGD平均降低/%",
    "igd_reduction_std_percent": "IGD降低标准差/%",
    "igd_improved_seeds": "IGD改善种子数",
    "sparsity_reduction_mean_percent": "稀疏度平均降低/%",
    "min_cost_change_mean_percent": "最低成本平均变化/%",
    "min_carbon_change_mean_percent": "最低碳排放平均变化/%",
    "front_points_change_mean": "帕累托点数平均变化",
    "cost_change_mean_percent": "成本平均变化/%",
    "cost_change_std_percent": "成本变化标准差/%",
    "carbon_change_mean_percent": "碳排放平均变化/%",
    "carbon_change_std_percent": "碳排放变化标准差/%",
    "grid_buy_change_mean_percent": "电网购电平均变化/%",
    "peer_trade_change_mean_percent": "园区交易平均变化/%",
    "waste_heat_change_mean_percent": "弃热平均变化/%",
    "paired_t_pvalue": "配对t检验p值",
    "wilcoxon_pvalue": "Wilcoxon p值",
    "mean_improvement_percent": "平均改善/%",
    "bootstrap_ci95_low": "95%区间下限",
    "bootstrap_ci95_high": "95%区间上限",
    "rank_biserial": "秩二列相关",
    "num_pairs": "配对数",
    "total_cost_mean": "成本均值",
    "total_cost_std": "成本标准差",
    "total_carbon_mean": "碳排放均值",
    "total_carbon_std": "碳排放标准差",
    "grid_buy_mean": "电网购电均值",
    "grid_buy_std": "电网购电标准差",
    "peer_trade_mean": "园区交易均值",
    "peer_trade_std": "园区交易标准差",
    "max_balance_error": "最大平衡误差",
    "level": "分解层级",
    "effect_unit": "效应单位",
    "alternative": "备择假设",
    "mean_effect": "平均效应",
    "std_effect": "效应标准差",
    "improved_seeds": "改善种子数",
    "total_cost": "运行成本",
    "total_carbon": "碳排放量",
    "total_grid_buy": "电网购电量",
    "total_peer_trade": "园区交易电量",
    "policy_file": "策略文件",
    "pareto_index": "帕累托索引",
    "cost": "运行成本",
    "carbon": "碳排放量",
}


VALUE_LABELS = {
    "spring": "春季",
    "summer": "夏季",
    "autumn": "秋季",
    "winter": "冬季",
    "cost_priority": "成本优先 [0.8, 0.2]",
    "balanced": "相对均衡 [0.5, 0.5]",
    "carbon_priority": "碳排放优先 [0.2, 0.8]",
    "成本优先": "成本优先 [0.8, 0.2]",
    "相对均衡": "相对均衡 [0.5, 0.5]",
    "均衡偏好": "相对均衡 [0.5, 0.5]",
    "碳排放优先": "碳排放优先 [0.2, 0.8]",
    "residential": "居民区",
    "commercial": "商业区",
    "industry": "工业区",
    "A0_full": "主方法",
    "A1_standard_cda": "A1 碳红利机制消融",
    "A3_random_preference": "A3 偏好选择机制消融",
    "A5_scalar_stale": "A5 向量经验回放消融",
    "normalized_hv": "归一化超体积",
    "raw_hypervolume": "原始超体积",
    "igd": "IGD",
    "normalized_igd": "归一化IGD",
    "normalized_sparsity": "归一化稀疏度",
    "front_points": "帕累托点数",
    "min_cost": "最低成本",
    "min_carbon": "最低碳排放",
    "greater": "大于",
    "less": "小于",
    "two-sided": "双侧",
    "preference": "偏好",
    "season": "季节",
    "cell": "偏好-季节组合",
}


def read_csv(relative_path: str) -> list[dict[str, str]]:
    path = RESULTS / Path(relative_path)
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def is_number(value: str) -> bool:
    try:
        float(value)
        return True
    except (TypeError, ValueError):
        return False


def format_value(value: object, column: str = "") -> str:
    if value is None:
        return "-"
    text = str(value).strip()
    if not text:
        return "-"
    if text in VALUE_LABELS:
        return VALUE_LABELS[text]
    if not is_number(text):
        return text

    number = float(text)
    if not math.isfinite(number):
        return text
    key = column.lower()
    if any(token in key for token in ("seed", "points", "wins", "losses", "ties", "排名", "种子数", "帕累托索引", "完成种子")):
        if abs(number - round(number)) < 1e-9:
            return str(int(round(number)))
    if "pvalue" in key or "p值" in column:
        return f"{number:.3e}"
    if any(token in key for token in ("weight", "权重", "占比", "率")):
        return f"{number:.4f}"
    if any(token in key for token in ("percent", "百分比", "变化/%", "改善/%")):
        return f"{number:.3f}"
    if abs(number) >= 1e9:
        return f"{number:.6e}"
    if abs(number) >= 1000:
        return f"{number:,.2f}"
    if number != 0 and abs(number) < 0.001:
        return f"{number:.3e}"
    return f"{number:.4f}"


def table_rows(
    relative_path: str,
    columns: Sequence[str],
    row_filter: Callable[[dict[str, str]], bool] | None = None,
) -> list[list[str]]:
    records = read_csv(relative_path)
    if row_filter:
        records = [record for record in records if row_filter(record)]
    return [[format_value(record.get(column, ""), column) for column in columns] for record in records]


def labels(columns: Sequence[str]) -> list[str]:
    return [COLUMN_LABELS.get(column, column.replace("_", " ")) for column in columns]


def set_cell_text(cell, text: str, *, bold: bool = False, size: float = 7.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_border(cell, **edges) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        settings = edges.get(edge, {"val": "nil"})
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in settings.items():
            element.set(qn(f"w:{key}"), str(value))


def apply_three_line_style(table) -> None:
    nil = {"val": "nil"}
    top = {"val": "single", "sz": "12", "color": "000000"}
    middle = {"val": "single", "sz": "6", "color": "000000"}
    bottom = {"val": "single", "sz": "12", "color": "000000"}
    last_index = len(table.rows) - 1
    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        for cell in row.cells:
            edges = {"top": nil, "bottom": nil, "left": nil, "right": nil, "insideH": nil, "insideV": nil}
            if row_index == 0:
                edges["top"] = top
                edges["bottom"] = middle
            if row_index == last_index:
                edges["bottom"] = bottom
            set_cell_border(cell, **edges)


def add_source(document: Document, relative_path: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(f"数据来源：results/{relative_path.replace(chr(92), '/')}")
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(7.5)
    run.font.italic = True
    run.font.color.rgb = None


def add_table(
    document: Document,
    number: int,
    title: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[str]],
    source: str,
) -> None:
    rows = list(rows)
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.space_before = Pt(7)
    caption.paragraph_format.space_after = Pt(4)
    run = caption.add_run(f"表 {number}  {title}")
    run.bold = True
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    run.font.size = Pt(10.5)

    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, text in enumerate(headers):
        set_cell_text(header.cells[index], text, bold=True, size=7.5)

    for row_values in rows:
        row = table.add_row()
        for index, text in enumerate(row_values):
            set_cell_text(row.cells[index], str(text), size=7.2)

    apply_three_line_style(table)
    add_source(document, source)


def add_csv_table(
    document: Document,
    number: int,
    title: str,
    path: str,
    columns: Sequence[str],
    row_filter: Callable[[dict[str, str]], bool] | None = None,
) -> None:
    add_table(document, number, title, labels(columns), table_rows(path, columns, row_filter), path)


def add_section(document: Document, title: str, intro: str) -> None:
    if len(document.paragraphs) > 3:
        document.add_page_break()
    heading = document.add_heading(title, level=1)
    heading.paragraph_format.space_after = Pt(5)
    paragraph = document.add_paragraph(intro)
    paragraph.paragraph_format.space_after = Pt(6)


def trade_hourly_summary() -> list[list[str]]:
    path = "04_园区电能交易结果/peer_trade_hourly_detail.csv"
    groups: dict[tuple[str, ...], list[float]] = defaultdict(list)
    for record in read_csv(path):
        key = (
            record["偏好类型"],
            record["典型日"],
            record["本园区"],
            record["交易对象园区"],
        )
        groups[key].append(float(record["交易电量_kWh_买入为正卖出为负"]))
    rows = []
    for key in sorted(groups):
        values = groups[key]
        rows.append(
            [
                *key,
                format_value(sum(values), "交易净电量"),
                format_value(sum(v for v in values if v > 0), "累计买入量"),
                format_value(-sum(v for v in values if v < 0), "累计卖出量"),
                format_value(max(values, default=0.0), "最大单时段买入量"),
                format_value(abs(min(values, default=0.0)), "最大单时段卖出量"),
            ]
        )
    return [[VALUE_LABELS.get(v, v) for v in row[:4]] + row[4:] for row in rows]


def trade_core_summary() -> list[list[str]]:
    """Aggregate the four typical days into a thesis-sized counterparty table."""
    path = "04_园区电能交易结果/peer_trade_hourly_detail.csv"
    groups: dict[tuple[str, ...], list[float]] = defaultdict(list)
    for record in read_csv(path):
        key = (record["偏好类型"], record["本园区"], record["交易对象园区"])
        groups[key].append(float(record["交易电量_kWh_买入为正卖出为负"]))
    rows = []
    for key in sorted(groups):
        values = groups[key]
        rows.append(
            [
                *(VALUE_LABELS.get(value, value) for value in key),
                format_value(sum(values), "四典型日净交易电量"),
                format_value(sum(value for value in values if value > 0), "累计买入电量"),
                format_value(-sum(value for value in values if value < 0), "累计卖出电量"),
                format_value(max(values, default=0.0), "最大单时段买入量"),
                format_value(abs(min(values, default=0.0)), "最大单时段卖出量"),
            ]
        )
    return rows


def balance_residual_summary() -> list[list[str]]:
    path = "03_主方法调度结果/调度图/balance_residuals.csv"
    groups: dict[tuple[str, str], dict[str, float]] = defaultdict(lambda: {"raw": 0.0, "plot": 0.0})
    for record in read_csv(path):
        key = (record["偏好类型"], record["能源类型"])
        groups[key]["raw"] = max(groups[key]["raw"], abs(float(record["原始最大绝对平衡残差"])))
        groups[key]["plot"] = max(groups[key]["plot"], abs(float(record["绘图后最大绝对平衡残差"])))
    return [
        [key[0], key[1], format_value(value["raw"]), format_value(value["plot"])]
        for key, value in sorted(groups.items())
    ]


def configure_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.25

    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    return document


def build_full_report() -> tuple[int, int]:
    document = configure_document()
    title = document.add_heading("实验结果数据表汇编", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(
        "多智能体进化强化学习综合能源系统优化调度项目\n"
        f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note = document.add_paragraph(
        "说明：本文档以 results 文件夹中保存的正式实验数据为唯一数据源。所有表格均采用三线表；"
        "成本、碳排放、电量等物理量保留合理有效位数，显著性检验 p 值采用科学计数法。"
        "逐小时和逐代原始明细保留在相应 CSV 文件中，本文档给出与实验图对应的汇总、机制和验证数据。"
    )
    note.paragraph_format.space_before = Pt(12)

    table_number = 0

    def csv_table(title: str, path: str, columns: Sequence[str], row_filter=None) -> None:
        nonlocal table_number
        table_number += 1
        add_csv_table(document, table_number, title, path, columns, row_filter)

    def custom_table(title: str, headers: Sequence[str], rows: Iterable[Sequence[str]], source: str) -> None:
        nonlocal table_number
        table_number += 1
        add_table(document, table_number, title, headers, rows, source)

    add_section(document, "一、全年真实数据跨算法对比", "对各算法在全年逐时真实数据上的帕累托优化质量、稳定性和统计显著性进行比较。")
    path = "01_全年真实数据跨算法对比/表格与说明/年度五随机种子汇总.csv"
    csv_table("全年跨算法五随机种子结果", path, ["方法", "全年原始超体积均值_CNY乘kgCO2", "全年原始超体积标准差_CNY乘kgCO2", "全年归一化IGD均值", "全年归一化IGD标准差", "全年归一化稀疏度均值", "全年帕累托点数均值"])
    path = "01_全年真实数据跨算法对比/表格与说明/年度合并前沿指标.csv"
    csv_table("全年跨算法合并帕累托前沿指标", path, ["方法", "随机种子数", "合并前沿帕累托点数", "合并前沿原始超体积_CNY乘kgCO2", "合并前沿归一化稀疏度", "合并前沿超体积排名", "合并前沿稀疏度排名"])
    path = "01_全年真实数据跨算法对比/表格与说明/年度配对显著性检验.csv"
    csv_table("主方法与对比算法的全年配对显著性检验", path, ["对比方法", "指标", "主方法胜出种子数", "主方法落后种子数", "主方法平均改善量", "单侧Wilcoxon检验p值"])

    add_section(document, "二、四季典型日跨算法对比", "对各算法在春、夏、秋、冬四个典型日组成的测试集上进行多目标性能比较。")
    path = "02_四季典型日跨算法对比/表格与说明/五随机种子汇总.csv"
    csv_table("典型日跨算法五随机种子结果", path, ["方法", "原始超体积均值_CNY乘kgCO2", "原始超体积标准差_CNY乘kgCO2", "归一化IGD均值", "归一化IGD标准差", "归一化稀疏度均值", "帕累托点数均值", "最低运行成本均值_CNY", "最低碳排放均值_kgCO2"])
    path = "02_四季典型日跨算法对比/表格与说明/pooled_front_metrics.csv"
    csv_table("典型日跨算法合并帕累托前沿指标", path, ["方法", "随机种子数", "合并前沿帕累托点数", "合并前沿原始超体积_CNY乘kgCO2", "合并前沿超体积排名", "合并前沿归一化稀疏度", "合并前沿稀疏度排名"])
    path = "02_四季典型日跨算法对比/表格与说明/配对显著性检验.csv"
    csv_table("典型日跨算法配对显著性检验", path, ["对比方法", "指标", "主方法胜出种子数", "主方法落后种子数", "主方法平均改善量", "单侧Wilcoxon检验p值"])

    add_section(document, "三、主方法典型调度结果", "列出三种典型偏好下的策略选择、系统及园区调度结果，并验证能量、交易资金和碳红利预算平衡。")
    path = "03_主方法调度结果/统计表/selected_policies.csv"
    csv_table("三种典型偏好对应的代表性策略", path, ["偏好类型", "策略文件", "帕累托索引", "目标成本权重", "目标碳排放权重", "成本权重", "碳排放权重", "权重距离", "系统运行成本", "系统碳排放量_kgCO2"])
    path = "03_主方法调度结果/统计表/preference_system_summary.csv"
    csv_table("三种典型偏好的系统级调度汇总", path, ["偏好类型", "成本权重", "碳排放权重", "运行成本", "碳排放量_kgCO2", "园区间成交电量_kWh", "电网购电量_kWh", "电网上网电量_kWh", "碳征费", "碳红利", "弃热量"])
    path = "03_主方法调度结果/统计表/preference_agent_summary.csv"
    csv_table("三种典型偏好的园区级调度汇总", path, ["偏好类型", "园区", "运行成本", "碳排放量_kgCO2", "园区间总购电量_kWh_买入为正卖出为负", "电网购电量_kWh", "电网上网电量_kWh", "碳征费", "碳红利", "弃热量"])
    path = "03_主方法调度结果/统计表/daily_system_summary.csv"
    csv_table("三种偏好和四季典型日的系统调度结果", path, ["偏好类型", "典型日", "运行成本", "碳排放量_kgCO2", "园区间成交电量_kWh", "电网购电量_kWh", "电网上网电量_kWh", "碳征费", "碳红利", "弃热量", "最大电功率平衡绝对残差"])
    path = "03_主方法调度结果/统计表/daily_agent_summary.csv"
    csv_table("三种偏好和四季典型日的分园区调度结果", path, ["偏好类型", "典型日", "园区", "运行成本", "碳排放量_kgCO2", "园区间总购电量_kWh_买入为正卖出为负", "电网购电量_kWh", "电网上网电量_kWh", "弃热量", "最大电功率平衡绝对残差"])
    path = "03_主方法调度结果/统计表/validation_detail.csv"
    csv_table("调度结果一致性与平衡约束验证", path, ["偏好类型", "典型日", "严格测试运行成本误差", "严格测试碳排放误差", "园区间电量平衡误差_kWh", "交易资金平衡误差", "碳红利预算平衡误差", "弃热量", "最大电功率平衡绝对残差"])
    custom_table("调度图能量平衡残差汇总", ["偏好类型", "能源类型", "原始最大绝对残差", "绘图后最大绝对残差"], balance_residual_summary(), "03_主方法调度结果/调度图/balance_residuals.csv")

    add_section(document, "四、园区间电能交易结果", "给出每个园区与两个交易对象之间的日电量和逐时交易统计；正值表示买入，负值表示卖出。")
    path = "04_园区电能交易结果/peer_trade_daily_by_counterparty.csv"
    csv_table("不同偏好和典型日下的园区间交易电量", path, ["偏好类型", "典型日", "本园区", "交易对象园区", "交易电量_kWh_买入为正卖出为负"])
    custom_table(
        "园区间逐时电能交易统计",
        ["偏好类型", "典型日", "本园区", "交易对象园区", "净交易电量/kWh", "累计买入/kWh", "累计卖出/kWh", "最大单时段买入/kWh", "最大单时段卖出/kWh"],
        trade_hourly_summary(),
        "04_园区电能交易结果/peer_trade_hourly_detail.csv",
    )

    add_section(document, "五、收敛性与多随机种子验证", "通过十随机种子重复实验比较主方法与 MADDPG 的最终性能和收敛过程。")
    path = "05_收敛性与多随机种子验证/主方法与MADDPG对比10种子/summary.csv"
    csv_table("主方法与MADDPG十随机种子性能汇总", path, ["method", "hv_mean", "hv_std", "igd_mean", "igd_std", "sparsity_mean", "front_points_mean"])
    path = "05_收敛性与多随机种子验证/主方法与MADDPG对比10种子/paired_tests.csv"
    csv_table("主方法与MADDPG十随机种子配对检验", path, ["metric", "wins", "ties", "losses", "mean_improvement", "bootstrap_ci_low", "bootstrap_ci_high", "one_sided_wilcoxon_p", "cohen_dz"])
    path = "05_收敛性与多随机种子验证/主方法收敛10种子/auc_summary.csv"
    csv_table("主方法与MADDPG收敛曲线面积汇总", path, ["method", "hv_auc_mean", "hv_auc_std", "igd_auc_mean", "igd_auc_std"])
    path = "05_收敛性与多随机种子验证/主方法收敛10种子/summary.csv"
    csv_table("主方法与MADDPG完整收敛过程数据", path, ["method", "evaluations", "hours_mean", "hv_mean", "hv_std", "igd_mean", "igd_std", "points_mean"])

    add_section(document, "六、主方法消融实验", "从全年性能和模块专属机制指标两个层面验证碳红利 CDA、偏好引导选择和向量经验回放的作用。")
    path = "06_消融实验/主方法全年消融/表格与协议/年度五随机种子汇总.csv"
    csv_table("主方法及A1、A3、A5消融方法的五随机种子结果", path, ["variant", "method", "hv_mean", "hv_std", "igd_mean", "igd_std", "sparsity_mean", "front_points_mean", "min_cost_mean", "min_carbon_mean"])
    path = "06_消融实验/主方法全年消融/表格与协议/年度合并前沿指标.csv"
    csv_table("主方法及A1、A3、A5消融方法的全年合并前沿指标", path, ["variant", "method", "seeds", "front_points", "raw_hypervolume", "normalized_igd", "normalized_sparsity", "min_annual_cost", "min_annual_carbon"])
    path = "06_消融实验/主方法全年消融/表格与协议/年度配对显著性检验.csv"
    csv_table("全年消融实验配对显著性检验", path, ["ablation", "ablation_method", "metric", "main_wins", "main_losses", "mean_improvement", "one_sided_wilcoxon_p"])
    base = "06_消融实验/主方法全年消融/模块作用验证/表格与协议"
    csv_table("三个核心模块的机制相对效应", f"{base}/模块机制相对效应汇总.csv", ["消融模块", "机制指标", "完整方法", "消融方法", "期望方向", "相对改善百分比"])
    csv_table("三个核心模块机制指标的配对检验", f"{base}/模块机制指标配对检验.csv", ["消融模块", "机制指标", "完整方法均值", "消融方法均值", "完整方法减消融方法", "双侧Wilcoxon检验p值"])
    csv_table("A1碳红利机制消融的作用指标", f"{base}/A1碳红利机制总体汇总.csv", ["方法", "全年碳排放均值_kgCO2", "园区交易电量均值_kWh", "电网购电量均值_kWh", "交易卖方平均碳强度", "最低碳卖方交易占比", "碳红利资金均值_CNY", "最大预算平衡误差_CNY", "碳红利生效小时占比"])
    csv_table("A3偏好选择机制消融的作用指标", f"{base}/A3偏好选择总体汇总.csv", ["方法", "平均有效偏好数量", "平均重复偏好率", "平均偏好覆盖宽度", "双极端覆盖率", "平均偏好切换幅度", "最终训练档案规模", "策略档案接受率"])
    csv_table("A5向量经验回放消融的奖励错配指标", f"{base}/A5奖励错配总体汇总.csv", ["方法", "历史经验偏好不一致占比", "潜在奖励相对错配率", "实际奖励相对错配率", "最大实际奖励相对错配率"])

    add_section(document, "七、补充消融与机制实验", "整理形成当前方法之前完成的碳红利机制、CDA与算法析因、精英选择以及经验回放实验。")
    base = "06_消融实验/Carbon_Dividend_V3机制消融"
    csv_table("Carbon-Dividend CDA V3五随机种子消融结果", f"{base}/v3_dividend_ablation_5seed_summary.csv", ["method", "completed_seeds", "hv_mean", "hv_std", "igd_mean", "igd_std", "sparsity_mean", "sparsity_std", "min_cost_mean", "min_carbon_mean", "front_points_mean"])
    csv_table("Carbon-Dividend CDA V3消融的配对效应", f"{base}/v3_dividend_ablation_5seed_paired_summary.csv", ["comparison", "hv_change_mean_percent", "hv_improved_seeds", "igd_reduction_mean_percent", "igd_improved_seeds", "sparsity_reduction_mean_percent", "min_cost_change_mean_percent", "min_carbon_change_mean_percent", "front_points_change_mean"])
    csv_table("Carbon-Dividend CDA V3调度行为效应", f"{base}/v3_dividend_ablation_behavior_effect_summary_5seed.csv", ["comparison", "cost_change_mean_percent", "carbon_change_mean_percent", "grid_buy_change_mean_percent", "peer_trade_change_mean_percent", "waste_heat_change_mean_percent"])
    base = "06_消融实验/CDA与进化机制析因消融/表格与协议"
    csv_table("CDA交易机制与进化算法析因实验", f"{base}/summary.csv", ["market", "algorithm", "hv_mean", "hv_std", "igd_mean", "igd_std", "sparsity_mean", "front_points_mean", "min_cost_mean", "min_carbon_mean"])
    csv_table("CDA交易机制与进化算法的主效应", f"{base}/factor_effects.csv", ["effect", "metric", "wins", "losses", "mean_improvement", "one_sided_p"])
    base = "06_消融实验/精英选择消融/表格与协议"
    csv_table("精英选择机制消融结果", f"{base}/summary.csv", ["method", "hv_mean", "hv_std", "igd_mean", "igd_std", "sparsity_mean", "sparsity_std", "front_points_mean"])
    csv_table("精英选择机制消融配对检验", f"{base}/paired_tests.csv", ["comparator", "metric", "wins", "losses", "mean_improvement", "one_sided_p"])
    base = "06_消融实验/向量经验回放消融/表格与协议"
    csv_table("向量经验回放机制消融结果", f"{base}/summary.csv", ["method", "hv_mean", "hv_std", "igd_mean", "igd_std", "sparsity_mean", "sparsity_std", "front_points_mean"])
    csv_table("向量经验回放机制消融配对检验", f"{base}/paired_tests.csv", ["comparator", "metric", "wins", "losses", "mean_improvement", "one_sided_wilcoxon_p"])

    add_section(document, "八、碳价敏感性分析", "比较不同碳价设置下主方法的帕累托优化质量、极值目标和解集规模。")
    base = "07_碳价敏感性分析/表格与协议"
    csv_table("不同碳价下的五随机种子性能", f"{base}/summary.csv", ["carbon_price", "hv_mean", "hv_std", "igd_mean", "igd_std", "sparsity_mean", "sparsity_std", "front_points_mean", "min_cost_mean", "min_carbon_mean"])
    csv_table("不同碳价与基准碳价的配对检验", f"{base}/paired_tests.csv", ["comparison", "metric", "wins", "losses", "mean_improvement", "one_sided_wilcoxon_p"])

    add_section(document, "九、偏好与季节分解分析", "从偏好、季节以及偏好与季节组合三个层面分析方法的运行成本、碳排放和能量交互行为。")
    base = "08_偏好与季节分解分析"
    csv_table("偏好与季节组合的五随机种子结果", f"{base}/v3_preference_season_5seed_cell_summary.csv", ["method", "preference", "season", "total_cost_mean", "total_cost_std", "total_carbon_mean", "total_carbon_std", "grid_buy_mean", "peer_trade_mean", "completed_seeds", "max_balance_error"])
    csv_table("按偏好汇总的五随机种子结果", f"{base}/v3_preference_season_5seed_preference_summary.csv", ["method", "preference", "total_cost_mean", "total_cost_std", "total_carbon_mean", "total_carbon_std", "grid_buy_mean", "peer_trade_mean", "completed_seeds", "max_balance_error"])
    csv_table("按季节汇总的五随机种子结果", f"{base}/v3_preference_season_5seed_season_summary.csv", ["method", "season", "total_cost_mean", "total_cost_std", "total_carbon_mean", "total_carbon_std", "grid_buy_mean", "peer_trade_mean", "completed_seeds", "max_balance_error"])
    csv_table("偏好层面的效应与显著性统计", f"{base}/v3_preference_season_5seed_preference_statistics.csv", ["comparison", "preference", "metric", "effect_unit", "mean_effect", "bootstrap_ci95_low", "bootstrap_ci95_high", "improved_seeds", "wilcoxon_pvalue"])
    csv_table("季节层面的效应与显著性统计", f"{base}/v3_preference_season_5seed_season_statistics.csv", ["comparison", "season", "metric", "effect_unit", "mean_effect", "bootstrap_ci95_low", "bootstrap_ci95_high", "improved_seeds", "wilcoxon_pvalue"])
    csv_table("偏好与季节组合的效应与显著性统计", f"{base}/v3_preference_season_5seed_cell_statistics.csv", ["comparison", "preference", "season", "metric", "effect_unit", "mean_effect", "improved_seeds", "wilcoxon_pvalue"])

    add_section(document, "十、主方法与基线的调度行为对比", "对主方法和 MADDPG 在三种偏好及四季典型日下的成本、碳排放、电网购电和园区交易行为进行比较。")
    base = "09_调度行为对比/表格"
    csv_table("主方法与MADDPG代表性策略", f"{base}/selected_policies.csv", ["method", "preference", "policy_file", "pareto_index", "cost", "carbon"])
    csv_table("主方法与MADDPG不同偏好下的调度行为", f"{base}/summary.csv", ["method", "preference", "total_cost", "total_carbon", "total_grid_buy", "total_peer_trade"])
    csv_table("主方法与MADDPG不同偏好和季节下的调度行为", f"{base}/season_summary.csv", ["method", "preference", "season", "total_cost", "total_carbon", "total_grid_buy", "total_peer_trade"])

    document.add_page_break()
    document.add_heading("附：表格与原始数据对应关系", level=1)
    paragraph = document.add_paragraph(
        "文档中的表格用于精确描述 results 中各实验图展示的结果。帕累托前沿全部点、逐随机种子、逐代、"
        "逐小时调度和逐小时交易明细因行数较多，继续保存在各实验目录的 CSV 文件中；表下方的数据来源可直接定位。"
    )
    paragraph.paragraph_format.space_after = Pt(8)
    custom_table(
        "实验类别与正文数据表覆盖关系",
        ["实验类别", "对应表号范围", "主要描述内容"],
        [
            ["全年跨算法对比", "表1-表3", "全年前沿质量、合并前沿和显著性"],
            ["典型日跨算法对比", "表4-表6", "四季典型日性能和显著性"],
            ["主方法调度结果", "表7-表13", "策略、系统、园区和约束验证"],
            ["园区电能交易", "表14-表15", "交易对象级日统计和逐时统计"],
            ["收敛性与多种子", "表16-表19", "十种子性能、检验和收敛过程"],
            ["当前主方法消融", "表20-表27", "A1/A3/A5总体性能和模块作用"],
            ["补充消融实验", "表28-表36", "CDA、精英选择和经验回放机制"],
            ["碳价敏感性", "表37-表38", "不同碳价的性能和检验"],
            ["偏好与季节分解", "表39-表44", "偏好、季节和组合效应"],
            ["调度行为对比", "表45-表47", "主方法相对MADDPG的调度行为"],
        ],
        "结果文件清单.csv",
    )

    document.core_properties.title = "实验结果数据表汇编"
    document.core_properties.subject = "多智能体进化强化学习综合能源系统优化调度实验结果"
    document.core_properties.author = "项目实验结果自动汇编"
    document.save(OUTPUT)
    return table_number, sum(1 for _ in RESULTS.rglob("*.csv"))


def build_report() -> tuple[int, int]:
    """Build the concise, thesis-facing collection of core result tables."""
    document = configure_document()
    title = document.add_heading("核心实验结果数据表", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(
        "多智能体进化强化学习综合能源系统优化调度项目\n"
        f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note = document.add_paragraph(
        "说明：本文档仅保留能够直接支撑论文主要结论的核心实验数据。偏好向量统一按照"
        "[成本权重, 碳排放权重] 的顺序书写；成本优先、相对均衡和碳排放优先分别对应"
        "[0.8, 0.2]、[0.5, 0.5] 和 [0.2, 0.8]。逐小时、逐代及全部帕累托点等原始明细"
        "仍保存在 results 对应目录的 CSV 文件中，不在论文表格中重复罗列。"
    )
    note.paragraph_format.space_before = Pt(12)

    table_number = 0

    def csv_table(title_text: str, path: str, columns: Sequence[str]) -> None:
        nonlocal table_number
        table_number += 1
        add_csv_table(document, table_number, title_text, path, columns)

    def custom_table(title_text: str, headers: Sequence[str], rows: Iterable[Sequence[str]], source: str) -> None:
        nonlocal table_number
        table_number += 1
        add_table(document, table_number, title_text, headers, rows, source)

    add_section(
        document,
        "一、跨算法优化性能",
        "使用全年8760小时真实数据作为主要比较依据，并以四季典型日结果验证算法在代表性负荷场景下的表现。",
    )
    path = "01_全年真实数据跨算法对比/表格与说明/年度五随机种子汇总.csv"
    csv_table(
        "全年跨算法五随机种子性能",
        path,
        ["方法", "全年原始超体积均值_CNY乘kgCO2", "全年原始超体积标准差_CNY乘kgCO2", "全年归一化IGD均值", "全年归一化IGD标准差", "全年归一化稀疏度均值", "全年帕累托点数均值"],
    )
    path = "01_全年真实数据跨算法对比/表格与说明/年度合并前沿指标.csv"
    csv_table(
        "全年跨算法合并帕累托前沿指标",
        path,
        ["方法", "随机种子数", "合并前沿帕累托点数", "合并前沿原始超体积_CNY乘kgCO2", "合并前沿归一化稀疏度", "合并前沿超体积排名", "合并前沿稀疏度排名"],
    )
    path = "01_全年真实数据跨算法对比/表格与说明/年度配对显著性检验.csv"
    csv_table(
        "主方法与对比算法的全年配对显著性检验",
        path,
        ["对比方法", "指标", "主方法胜出种子数", "主方法落后种子数", "主方法平均改善量", "单侧Wilcoxon检验p值"],
    )
    path = "02_四季典型日跨算法对比/表格与说明/五随机种子汇总.csv"
    csv_table(
        "四季典型日跨算法五随机种子性能",
        path,
        ["方法", "原始超体积均值_CNY乘kgCO2", "原始超体积标准差_CNY乘kgCO2", "归一化IGD均值", "归一化IGD标准差", "归一化稀疏度均值", "帕累托点数均值", "最低运行成本均值_CNY", "最低碳排放均值_kgCO2"],
    )

    add_section(
        document,
        "二、主方法调度与园区交易",
        "选取三种标准偏好的代表性非支配策略，比较四季典型日和三类园区的成本、碳排放及电能交互，并验证约束满足情况。",
    )
    path = "03_主方法调度结果/统计表/selected_policies.csv"
    csv_table(
        "三种标准偏好的代表性策略",
        path,
        ["偏好类型", "策略文件", "帕累托索引", "目标成本权重", "目标碳排放权重", "成本权重", "碳排放权重", "权重距离", "系统运行成本", "系统碳排放量_kgCO2"],
    )
    path = "03_主方法调度结果/统计表/daily_system_summary.csv"
    csv_table(
        "三种偏好和四季典型日的系统调度结果",
        path,
        ["偏好类型", "典型日", "运行成本", "碳排放量_kgCO2", "园区间成交电量_kWh", "电网购电量_kWh", "电网上网电量_kWh", "碳征费", "碳红利", "弃热量", "最大电功率平衡绝对残差"],
    )
    path = "03_主方法调度结果/统计表/daily_agent_summary.csv"
    csv_table(
        "三种偏好和四季典型日的分园区调度结果",
        path,
        ["偏好类型", "典型日", "园区", "运行成本", "碳排放量_kgCO2", "园区间总购电量_kWh_买入为正卖出为负", "电网购电量_kWh", "电网上网电量_kWh", "弃热量", "最大电功率平衡绝对残差"],
    )
    custom_table(
        "四个典型日的分交易对象电量汇总",
        ["偏好向量", "本园区", "交易对象园区", "净交易电量/kWh", "累计买入/kWh", "累计卖出/kWh", "最大单时段买入/kWh", "最大单时段卖出/kWh"],
        trade_core_summary(),
        "04_园区电能交易结果/peer_trade_hourly_detail.csv",
    )
    path = "03_主方法调度结果/统计表/validation_detail.csv"
    csv_table(
        "典型调度结果的约束与一致性验证",
        path,
        ["偏好类型", "典型日", "严格测试运行成本误差", "严格测试碳排放误差", "园区间电量平衡误差_kWh", "交易资金平衡误差", "碳红利预算平衡误差", "弃热量", "最大电功率平衡绝对残差"],
    )

    add_section(
        document,
        "三、稳定性与收敛性",
        "使用十个随机种子检验主方法相对MADDPG的稳定优势，并以收敛曲线面积衡量整个训练过程的样本效率。",
    )
    path = "05_收敛性与多随机种子验证/主方法与MADDPG对比10种子/summary.csv"
    csv_table("主方法与MADDPG十随机种子性能", path, ["method", "hv_mean", "hv_std", "igd_mean", "igd_std", "sparsity_mean", "front_points_mean"])
    path = "05_收敛性与多随机种子验证/主方法与MADDPG对比10种子/paired_tests.csv"
    csv_table("主方法与MADDPG十随机种子配对检验", path, ["metric", "wins", "ties", "losses", "mean_improvement", "bootstrap_ci_low", "bootstrap_ci_high", "one_sided_wilcoxon_p", "cohen_dz"])
    path = "05_收敛性与多随机种子验证/主方法收敛10种子/auc_summary.csv"
    csv_table("主方法与MADDPG收敛曲线面积", path, ["method", "hv_auc_mean", "hv_auc_std", "igd_auc_mean", "igd_auc_std"])

    add_section(
        document,
        "四、主方法核心消融",
        "通过A1、A3和A5三种消融分别验证Carbon-Dividend CDA V3、全任务偏好进化选择和向量经验回放的贡献。",
    )
    path = "06_消融实验/主方法全年消融/表格与协议/年度五随机种子汇总.csv"
    csv_table(
        "主方法与三种消融方法的五随机种子性能",
        path,
        ["variant", "method", "hv_mean", "hv_std", "igd_mean", "igd_std", "sparsity_mean", "front_points_mean", "min_cost_mean", "min_carbon_mean"],
    )
    path = "06_消融实验/主方法全年消融/表格与协议/年度合并前沿指标.csv"
    csv_table(
        "主方法与三种消融方法的全年合并前沿",
        path,
        ["variant", "method", "seeds", "front_points", "raw_hypervolume", "normalized_igd", "normalized_sparsity", "min_annual_cost", "min_annual_carbon"],
    )
    path = "06_消融实验/主方法全年消融/表格与协议/年度配对显著性检验.csv"
    csv_table("全年消融实验配对显著性检验", path, ["ablation", "ablation_method", "metric", "main_wins", "main_losses", "mean_improvement", "one_sided_wilcoxon_p"])
    base = "06_消融实验/主方法全年消融/模块作用验证/表格与协议"
    csv_table("三个核心模块的机制相对效应", f"{base}/模块机制相对效应汇总.csv", ["消融模块", "机制指标", "完整方法", "消融方法", "期望方向", "相对改善百分比"])
    csv_table("三个核心模块机制指标的配对检验", f"{base}/模块机制指标配对检验.csv", ["消融模块", "机制指标", "完整方法均值", "消融方法均值", "完整方法减消融方法", "双侧Wilcoxon检验p值"])
    csv_table("碳红利CDA的机制作用", f"{base}/A1碳红利机制总体汇总.csv", ["方法", "全年碳排放均值_kgCO2", "园区交易电量均值_kWh", "电网购电量均值_kWh", "交易卖方平均碳强度", "最低碳卖方交易占比", "碳红利资金均值_CNY", "最大预算平衡误差_CNY", "碳红利生效小时占比"])
    csv_table("全任务偏好进化选择的机制作用", f"{base}/A3偏好选择总体汇总.csv", ["方法", "平均有效偏好数量", "平均重复偏好率", "平均偏好覆盖宽度", "双极端覆盖率", "平均偏好切换幅度", "最终训练档案规模", "策略档案接受率"])
    csv_table("向量经验回放的奖励一致性作用", f"{base}/A5奖励错配总体汇总.csv", ["方法", "历史经验偏好不一致占比", "潜在奖励相对错配率", "实际奖励相对错配率", "最大实际奖励相对错配率"])

    add_section(
        document,
        "五、参数敏感性及场景分解",
        "分析碳价变化对主方法的影响，并从偏好和季节两个维度说明方法的适应性及运行规律。",
    )
    path = "07_碳价敏感性分析/表格与协议/summary.csv"
    csv_table("不同碳价下的五随机种子性能", path, ["carbon_price", "hv_mean", "hv_std", "igd_mean", "igd_std", "sparsity_mean", "sparsity_std", "front_points_mean", "min_cost_mean", "min_carbon_mean"])
    base = "08_偏好与季节分解分析"
    csv_table("不同偏好向量下的五随机种子结果", f"{base}/v3_preference_season_5seed_preference_summary.csv", ["method", "preference", "total_cost_mean", "total_cost_std", "total_carbon_mean", "total_carbon_std", "grid_buy_mean", "peer_trade_mean", "completed_seeds", "max_balance_error"])
    csv_table("不同季节下的五随机种子结果", f"{base}/v3_preference_season_5seed_season_summary.csv", ["method", "season", "total_cost_mean", "total_cost_std", "total_carbon_mean", "total_carbon_std", "grid_buy_mean", "peer_trade_mean", "completed_seeds", "max_balance_error"])

    add_section(
        document,
        "六、主方法与基线的调度行为",
        "对主方法和MADDPG在三种标准偏好及四个典型日下的成本、碳排放、电网购电和园区交易进行直接比较。",
    )
    base = "09_调度行为对比/表格"
    csv_table("主方法与MADDPG不同偏好向量下的调度行为", f"{base}/summary.csv", ["method", "preference", "total_cost", "total_carbon", "total_grid_buy", "total_peer_trade"])
    csv_table("主方法与MADDPG不同偏好和季节下的调度行为", f"{base}/season_summary.csv", ["method", "preference", "season", "total_cost", "total_carbon", "total_grid_buy", "total_peer_trade"])

    document.core_properties.title = "核心实验结果数据表"
    document.core_properties.subject = "多智能体进化强化学习综合能源系统优化调度核心实验结果"
    document.core_properties.author = "项目实验结果自动汇编"
    document.save(OUTPUT)
    return table_number, sum(1 for _ in RESULTS.rglob("*.csv"))


if __name__ == "__main__":
    count, csv_count = build_report()
    print(f"Generated: {OUTPUT}")
    print(f"Tables: {count}")
    print(f"CSV files indexed in results: {csv_count}")
